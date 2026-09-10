#!/usr/bin/env python3
"""Build a two-tab workbook the reviewer can drop into the live Google Sheet.

The Drive connector this session has can read Drive files and create new ones; it cannot write
cells or add tabs to a native Google Sheet. So the safe path is an IMPORT file: Google Sheets'
File > Import > "Insert new sheet(s)" adds these two tabs and touches nothing that is already
there. Columns match the live NHAA tab exactly, including the Scope column.
"""
import json, os, sys, collections
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(HERE, "..", "..", "..", ".."))
AM = os.path.join(REPO, "docs", "qc", "portals", "smile-admin", "audit-master.json")
OUT = os.path.join(HERE, "out", "SMILE-Beggary-QC-sheets-to-import.xlsx")
HEAD = ["ID", "Screen", "Category", "Severity", "Issue (Design → Built)", "Recommended Fix (Dev)",
        "Figma URL", "Live URL", "Status", "Assignee", "Date", "Notes", "Scope"]
COVHEAD = ["Screen", "Section", "Figma URL", "Has Design?", "Built?", "QC Status", "# Findings", "Notes"]
NAVY = "FF003366"


def header(ws, head, widths):
    ws.append(head)
    for i, h in enumerate(head, 1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=NAVY)
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    ws.freeze_panes = "A2"


def main():
    am = json.load(open(AM))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SMILE Beggary"
    header(ws, HEAD, [22, 34, 22, 10, 86, 70, 34, 34, 10, 12, 12, 36, 10])
    for s in am["screens"]:
        for f in s["findings"]:
            ws.append([f["id"], s["name"], f["axis"], f["severity"],
                       f"Figma: {f['figma']}  →  Live: {f['live']}", f["fix"],
                       s.get("figmaUrl"), s.get("liveUrl"), "Open", None, None,
                       "env: dev" + ("  ·  Scope: Global — fix once, lands everywhere"
                                     if f.get("scope") == "Global" else ""),
                       f.get("scope", "Screen")])
    for r in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for c in r:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        r[0].font = Font(size=9, bold=True)
    ws.auto_filter.ref = f"A1:M{ws.max_row}"

    cov = wb.create_sheet("Coverage – SMILE Beggary")
    header(cov, COVHEAD, [44, 18, 34, 12, 10, 14, 12, 52])
    rows = json.load(open(os.path.join(HERE, "sheet", "all_rows.json")))
    per = collections.Counter()
    for s in am["screens"]:
        per[s["name"]] += len(s["findings"])
    for r in rows:
        cov.append([f"{r['role']} — {r['title']}", r["role"], r["figmaUrl"], "Yes", "Yes",
                    "In Review", per.get(r["title"], 0), None])
    cov.append(["Surveyor detail (/surveyors/<id>)", "super-admin", None, "No", "Yes", "Not compared", 0,
                "No design frame on 'Smile Beggary (Synced)' — the nearest frame is the surveyor LIST."])
    for label in ("View Catalog", "Edit Permissions", "Create Survey Location", "View Beneficiary",
                  "Dashboard chart tabs", "Add District"):
        cov.append([label, "super-admin", None, "Yes", "Yes", "Not reached", 0,
                    "The control resolves and the click lands, but the view does not finish loading "
                    "inside the capture window."])
    for label in ("State Nodal Officer", "Nodal Officer", "Implementing Agency"):
        cov.append([f"Role — {label}", label, None, "Yes", "Yes", "Skipped", 0,
                    "Skipped at the reviewer's instruction for this pass."])
    for r in cov.iter_rows(min_row=2, max_row=cov.max_row):
        for c in r:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    cov.auto_filter.ref = f"A1:H{cov.max_row}"

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    counts = collections.Counter(f["severity"] for s in am["screens"] for f in s["findings"])
    total = sum(counts.values())
    print(f"{OUT}")
    print(f"  'SMILE Beggary'            {ws.max_row - 1} rows")
    print(f"  'Coverage – SMILE Beggary' {cov.max_row - 1} rows")
    print(f"  Rollup row to paste: SMILE Beggary | {total} | {counts['Blocker']} | "
          f"{counts['Major']} | {counts['Minor']} | {counts['Nit']} | {total} | 0 | 0")


if __name__ == "__main__":
    sys.exit(main())
